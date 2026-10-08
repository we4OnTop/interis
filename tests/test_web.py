import json

import pytest
from fastapi.testclient import TestClient

from interis.analysis.analyze import AnalysisOptions, analyze
from interis.analysis.guide import parse_guide
from interis.config import Paths
from interis.web.app import create_app
from interis.web.store import Store
from tests.test_analysis import FakeEncoder, _interview

GUIDE = ("- F1: Wie sieht Ihr Arbeitsalltag aus?\n- F2: Welche Rolle spielt KI?\n"
         "- F3: Wie gehen Sie mit vertraulichen Daten um?\n- F4: Was wünschen Sie sich?\n")
BASE = "http://127.0.0.1:8765"
TOKEN = "test-token"  # noqa: S105 – fixed token for tests only
H = {"X-Interis": "1"}


@pytest.fixture
def data_dir(tmp_path):
    paths = Paths(tmp_path)
    paths.ensure()
    (tmp_path / "leitfaden.md").write_text(GUIDE, encoding="utf-8")
    t = _interview()
    t.meta.update({"interview_id": "T1", "audio": {"duration_s": 30.0, "sha256": "x"}})
    t.analysis = analyze(t, AnalysisOptions(guide=parse_guide(GUIDE), match_threshold=0.9,
                                            answer_z=0.5), None, FakeEncoder)
    (tmp_path / "exports" / "T1").mkdir(parents=True)
    (tmp_path / "exports" / "T1" / "T1.json").write_text(json.dumps(t.to_dict()),
                                                          encoding="utf-8")
    return paths


@pytest.fixture
def client(data_dir):
    c = TestClient(create_app(data_dir, TOKEN, 8765), base_url=BASE)
    assert c.post("/api/login", json={"token": TOKEN}, headers=H).status_code == 200
    return c


# ------------------------------------------------------------------ security

def test_foreign_host_header_is_rejected(data_dir):
    c = TestClient(create_app(data_dir, TOKEN, 8765), base_url="http://evil.example")
    assert c.get("/").status_code == 400


def test_api_requires_login(data_dir):
    c = TestClient(create_app(data_dir, TOKEN, 8765), base_url=BASE)
    assert c.get("/api/projects").status_code == 401
    assert c.post("/api/login", json={"token": "wrong"}, headers=H).status_code == 401
    assert c.get("/api/projects").status_code == 401


def test_session_cookie_is_httponly_and_strict(data_dir):
    c = TestClient(create_app(data_dir, TOKEN, 8765), base_url=BASE)
    r = c.post("/api/login", json={"token": TOKEN}, headers=H)
    cookie = r.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=strict" in cookie


def test_state_change_needs_custom_header_and_same_origin(client):
    body = {"interview": "T1", "turn": 3, "first": 0, "last": 2, "guide_code": "F3"}
    assert client.post("/api/links", json=body).status_code == 403
    assert client.post("/api/links", json=body,
                       headers={**H, "Origin": "http://evil.example"}).status_code == 403
    assert client.post("/api/links", json=body,
                       headers={**H, "Origin": BASE}).status_code == 200


def test_security_headers_and_no_api_docs(client):
    r = client.get("/")
    assert "default-src 'self'" in r.headers["content-security-policy"]
    assert r.headers["x-content-type-options"] == "nosniff"
    for path in ("/docs", "/redoc", "/openapi.json"):
        assert client.get(path).status_code == 404


def test_input_validation(client):
    bad_turn = {"interview": "T1", "turn": 99, "first": 0, "last": 0, "guide_code": "F1"}
    assert client.post("/api/links", json=bad_turn, headers=H).status_code == 422
    bad_code = {"interview": "T1", "turn": 1, "first": 0, "last": 0, "guide_code": "X9"}
    assert client.post("/api/links", json=bad_code, headers=H).status_code == 422
    assert client.get("/api/interviews/nope").status_code == 404


# ------------------------------------------------------------------ review logic

def _cell(client, code):
    return client.get("/api/projects/1/compare").json()["cells"]["T1"][code]


def test_compare_shows_asked_questions_with_dialogue(client):
    c = _cell(client, "F2")
    assert c["status"] == "asked"
    turns = c["exchanges"][0]["dialogue"]
    assert turns[0]["role"] == "interviewer" and turns[1]["role"] == "interviewee"
    assert any(p.get("question") and p["code"] == "F2" for p in turns[0]["pieces"])


def test_link_answer_to_unasked_question_marks_it_omitted(client):
    assert _cell(client, "F4")["status"] == "missing"
    r = client.post("/api/links", headers=H, json={
        "interview": "T1", "turn": 1, "first": 0, "last": 2, "guide_code": "F4",
        "omitted": True})
    c = _cell(client, "F4")
    assert c["status"] == "omitted"
    link = c["links"][0]
    assert link["type"] == "unasked" and link["from_code"] == "F1"

    client.patch(f"/api/links/{r.json()['id']}", headers=H, json={"omitted": False})
    assert _cell(client, "F4")["status"] == "answered_elsewhere"
    client.delete(f"/api/links/{r.json()['id']}", headers=H)
    assert _cell(client, "F4")["status"] == "missing"


def test_rejecting_a_suggestion_hides_it(client):
    sug = _cell(client, "F3")["suggestions"]
    assert sug, "fixture should produce a suggestion for F3"
    s = sug[0]
    client.post("/api/links", headers=H, json={
        "interview": "T1", "turn": s["turn"], "first": s["first"], "last": s["last"],
        "guide_code": "F3", "status": "rejected", "source": "suggestion"})
    assert not _cell(client, "F3")["suggestions"]


def test_reassigning_a_question(client):
    q = next(q for q in client.get("/api/interviews/T1").json()["questions"]
             if q["guide_code"] == "F2")
    client.post("/api/questions", headers=H, json={
        "interview": "T1", "turn": q["turn"], "first": q["first"], "last": q["last"],
        "guide_code": "F4", "match": "main"})
    assert _cell(client, "F2")["status"] == "missing"
    assert _cell(client, "F4")["status"] == "asked"

    client.post("/api/questions/reset", headers=H,
                json={"interview": "T1", "turn": q["turn"], "first": q["first"]})
    assert _cell(client, "F2")["status"] == "asked"


def test_rejecting_a_question(client):
    q = client.get("/api/interviews/T1").json()["questions"][0]
    client.post("/api/questions", headers=H, json={
        "interview": "T1", "turn": q["turn"], "first": q["first"], "last": q["last"],
        "match": "followup", "status": "rejected"})
    ids = [(x["turn"], x["first"]) for x in client.get("/api/interviews/T1").json()["questions"]]
    assert (q["turn"], q["first"]) not in ids


def test_audio_streaming_with_range(client, data_dir, tmp_path):
    assert client.get("/api/interviews/T1/audio").status_code == 404
    audio = tmp_path / "a.wav"
    audio.write_bytes(b"RIFF" + bytes(1000))
    Store(data_dir.root / "interis.db").register_interview("T1", audio)
    r = client.get("/api/interviews/T1/audio", headers={"Range": "bytes=0-9"})
    assert r.status_code == 206 and len(r.content) == 10


# ------------------------------------------------------------------ projects, upload, jobs

def test_existing_interviews_are_adopted_into_a_project(client):
    projects = client.get("/api/projects").json()
    assert [p["name"] for p in projects] == ["Bestehende Interviews"]
    assert projects[0]["questions"] == 4 and projects[0]["transcribed"] == 1
    p = client.get("/api/projects/1").json()
    assert p["interviews"][0]["id"] == "T1" and p["next_id"] == "I01"


def test_create_project_and_save_guide(client):
    pid = client.post("/api/projects", headers=H, json={"name": "Neu"}).json()["id"]
    assert client.get(f"/api/projects/{pid}").json()["guide"] is None
    r = client.put(f"/api/projects/{pid}/guide", headers=H, json={"text": "nur Text"})
    assert r.status_code == 422
    r = client.put(f"/api/projects/{pid}/guide", headers=H, json={"text": GUIDE})
    assert r.json() == {"questions": 4, "reanalyze": 0}
    assert client.get(f"/api/projects/{pid}").json()["guide_text"] == GUIDE


def test_guide_change_queues_reanalysis(client, data_dir):
    client.put("/api/projects/1/guide", headers=H, json={"text": GUIDE + "- F5: Noch was?\n"})
    jobs = Store(data_dir.root / "interis.db").jobs()
    assert [(j["kind"], j["interview_id"], j["status"]) for j in jobs] == \
        [("analyze", "T1", "queued")]
    # no duplicate while one is pending
    assert client.post("/api/projects/1/reanalyze", headers=H).json() == {"reanalyze": 0}


def _new_interview(client, iid="I01", parts=(b"\x00" * 5000,), ext="m4a"):
    assert client.post("/api/projects/1/interviews", headers=H,
                       json={"interview": iid}).status_code == 200
    for content in parts:
        r = client.post(f"/api/interviews/{iid}/parts?ext={ext}", headers=H, content=content)
        assert r.status_code == 200, r.text


def test_upload_parts_and_start_transcription(client, data_dir):
    _new_interview(client, parts=(b"a" * 5000, b"b" * 300))
    store = Store(data_dir.root / "interis.db")
    assert store.part_paths("I01") == [data_dir.audio / "I01-1.m4a", data_dir.audio / "I01-2.m4a"]
    assert (data_dir.audio / "I01-1.m4a").stat().st_size == 5000
    assert store.project_of("I01") == 1 and store.jobs(["I01"]) == []  # draft, not started

    row = next(i for i in client.get("/api/projects/1").json()["interviews"] if i["id"] == "I01")
    assert [p["size"] for p in row["parts"]] == [5000, 300] and not row["transcribed"]

    # reorder: second recording first
    assert client.put("/api/interviews/I01/parts", headers=H, json={"order": [1, 0]}).json()
    assert store.part_paths("I01")[0].name == "I01-2.m4a"
    assert client.put("/api/interviews/I01/parts", headers=H,
                      json={"order": [0]}).status_code == 422

    r = client.post("/api/interviews/I01/transcribe", headers=H,
                    json={"settings": {"model": "whisper-large-v3-turbo"}, "preset": "P"})
    job = store.job(r.json()["job"])
    assert job["kind"] == "transcribe"
    assert job["options"]["model"] == "whisper-large-v3-turbo"
    assert job["options"]["preset"] == "P" and job["options"]["speakers"] == 2
    # no changes while the job is pending
    assert client.post("/api/interviews/I01/parts?ext=wav", headers=H,
                       content=b"x").status_code == 409
    assert client.delete("/api/interviews/I01/parts/0", headers=H).status_code == 409


def test_remove_part_deletes_only_the_upload(client, data_dir):
    _new_interview(client, parts=(b"a", b"b"))
    assert client.delete("/api/interviews/I01/parts/0", headers=H).status_code == 200
    assert not (data_dir.audio / "I01-1.m4a").exists()
    assert Store(data_dir.root / "interis.db").part_paths("I01") == [data_dir.audio / "I01-2.m4a"]


def test_start_needs_a_recording(client):
    client.post("/api/projects/1/interviews", headers=H, json={"interview": "I01"})
    assert client.post("/api/interviews/I01/transcribe", headers=H,
                       json={}).status_code == 422


def test_retranscription_requires_confirmation_when_markings_exist(client, data_dir):
    store = Store(data_dir.root / "interis.db")
    audio = data_dir.root / "orig.wav"
    audio.write_bytes(b"RIFF")
    store.register_interview("T1", audio)
    client.post("/api/links", headers=H, json={"interview": "T1", "turn": 1, "first": 0,
                                               "last": 2, "guide_code": "F4"})
    assert client.post("/api/interviews/T1/transcribe", headers=H,
                       json={}).status_code == 409
    assert client.post("/api/interviews/T1/transcribe", headers=H,
                       json={"discard_markings": True}).status_code == 200
    # kept until the new transcript exists (the finished job removes them)
    assert len(store.links("T1")) == 1


@pytest.mark.parametrize("iid, ext, status", [
    ("I01", "exe", 422),
    ("I01", "wav", 200),
])
def test_part_extension_allow_list(client, iid, ext, status):
    client.post("/api/projects/1/interviews", headers=H, json={"interview": iid})
    r = client.post(f"/api/interviews/{iid}/parts?ext={ext}", headers=H, content=b"x")
    assert r.status_code == status


@pytest.mark.parametrize("iid, status", [("../x", 422), ("T1", 409), ("a b", 422)])
def test_interview_id_rules(client, iid, status):
    r = client.post("/api/projects/1/interviews", headers=H, json={"interview": iid})
    assert r.status_code == status


def test_parts_upload_needs_known_interview_and_csrf_header(client):
    assert client.post("/api/interviews/nope/parts?ext=wav", headers=H,
                       content=b"x").status_code == 404
    client.post("/api/projects/1/interviews", headers=H, json={"interview": "I01"})
    assert client.post("/api/interviews/I01/parts?ext=wav", content=b"x").status_code == 403


def test_docx_guide_import(client):
    import io

    from docx import Document

    doc = Document()
    doc.add_heading("Einstieg", level=2)
    doc.add_paragraph("Vorab ein paar Worte zum Ablauf.")
    doc.add_paragraph("Wie sieht Ihr Arbeitsalltag aus?")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Welche Rolle spielt KI?"
    buf = io.BytesIO()
    doc.save(buf)
    r = client.post("/api/guide/import-docx", headers=H, content=buf.getvalue())
    text = r.json()["text"]
    assert "## Einstieg" in text and "- Wie sieht Ihr Arbeitsalltag aus?" in text
    assert "- Welche Rolle spielt KI?" in text and "- Vorab" not in text
    assert client.post("/api/guide/import-docx", headers=H, content=b"nope").status_code == 422


def test_job_runner_runs_cancels_and_reports(data_dir, tmp_path):
    import sys
    import time

    from interis.web.jobs import JobRunner

    store = Store(data_dir.root / "interis.db")
    script = ("import json,time\n"
              "print(json.dumps({'stage': 'transcribe', 'fraction': 1.0}), flush=True)\n"
              "print('fertig', flush=True)\n")

    class Fake(JobRunner):
        def command(self, job):
            if job["options"].get("fail"):
                return [sys.executable, "-c", "print('ERROR: kaputt'); raise SystemExit(1)"]
            if job["options"].get("slow"):
                return [sys.executable, "-c", "import time; time.sleep(30)"]
            return [sys.executable, "-c", script]

    runner = Fake(data_dir, store, lambda _i: None, lambda _i: "")
    ok = store.add_job("transcribe", "A")
    bad = store.add_job("transcribe", "B", {"fail": True})
    slow = store.add_job("transcribe", "C", {"slow": True})
    runner.start()
    try:
        deadline = time.time() + 30
        while store.job(slow)["status"] != "running" and time.time() < deadline:
            time.sleep(0.1)
        assert runner.cancel(slow)
        while store.job(slow)["status"] == "running" and time.time() < deadline:
            time.sleep(0.1)
    finally:
        runner.stop()
    assert store.job(ok)["status"] == "done" and store.job(ok)["message"] == "fertig"
    assert store.job(ok)["stage"] == "transcribe"
    assert store.job(bad)["status"] == "failed" and "kaputt" in store.job(bad)["message"]
    assert store.job(slow)["status"] == "cancelled"


def test_job_command_uses_project_guide_and_hotwords(data_dir):
    import sys

    from interis.web.jobs import JobRunner

    store = Store(data_dir.root / "interis.db")
    pid = store.create_project("P", "Müller SAP")
    audio = data_dir.root / "a.wav"
    audio.write_bytes(b"RIFF")
    store.add_interview("I01", pid, audio)
    store.add_part("I01", audio)
    guide = data_dir.root / "g.md"
    runner = JobRunner(data_dir, store, lambda _i: guide, lambda _i: "Müller SAP")
    cmd = runner.command(store.job(store.add_job("transcribe", "I01",
                                                 {"model": "whisper-large-v3-turbo"})))
    assert cmd[cmd.index("transcribe") + 1:cmd.index("--progress-json")] == [str(audio)] * 2
    assert cmd[:4] == [sys.executable, "-I", "-m", "interis.cli"]  # isolated child
    assert cmd[cmd.index("--id") + 1] == "I01"
    assert cmd[cmd.index("--guide") + 1] == str(guide)
    assert cmd[cmd.index("--hotwords") + 1] == "Müller SAP"
    assert cmd[cmd.index("--model") + 1] == "whisper-large-v3-turbo"
    assert "--progress-json" in cmd and "--room-mic" not in cmd
    cmd = runner.command(store.job(store.add_job("transcribe", "I01", {"room_mic": True})))
    assert "--room-mic" in cmd


def test_old_database_gets_project_column(tmp_path):
    import sqlite3

    db = tmp_path / "interis.db"
    with sqlite3.connect(db) as c:
        c.execute("CREATE TABLE interviews (id TEXT PRIMARY KEY, audio_path TEXT, "
                  "added_at TEXT NOT NULL)")
        c.execute("INSERT INTO interviews VALUES ('X', NULL, 'now')")
    store = Store(db)
    assert store.project_of("X") is None
    store.assign_project(["X"], store.create_project("P"))
    assert store.project_of("X") == 1


def test_delete_interview_removes_files_and_decisions(client, data_dir):
    _new_interview(client, parts=(b"abc", b"def"), ext="wav")
    client.post("/api/interviews/I01/transcribe", headers=H, json={})
    store = Store(data_dir.root / "interis.db")
    job = store.jobs(["I01"])[0]["id"]
    assert client.delete("/api/interviews/I01", headers=H).status_code == 409  # job pending
    client.post(f"/api/jobs/{job}/cancel", headers=H)
    assert client.delete("/api/interviews/I01", headers=H).status_code == 200
    assert not any(data_dir.audio.glob("I01*")) and store.project_of("I01") is None

    outside = data_dir.root / "original.wav"
    outside.write_bytes(b"RIFF")
    store.register_interview("T1", outside)
    client.post("/api/links", headers=H, json={"interview": "T1", "turn": 1, "first": 0,
                                               "last": 2, "guide_code": "F4"})
    assert client.delete("/api/interviews/T1", headers=H).status_code == 200
    assert outside.exists(), "recordings outside the upload folder are never deleted"
    assert not (data_dir.exports / "T1").exists() and store.links("T1") == []
    assert client.delete("/api/interviews/..", headers=H).status_code in (404, 405)


def test_exchange_shows_question_turn_even_if_next_question_is_in_same_turn():
    from interis.web.review import interview_state
    from tests.test_analysis import _interview

    t = _interview()
    t.analysis = analyze(t, AnalysisOptions(guide=parse_guide(GUIDE), match_threshold=0.9,
                                            answer_z=0.5), None, FakeEncoder)
    qs = [q for q in t.analysis["questions"] if q["match"] == "main"]
    # simulate a diarization error: two guide questions inside one turn
    second = {**qs[1], "turn": qs[0]["turn"], "first": qs[0]["last"] + 1,
              "last": qs[0]["last"] + 1}
    t.analysis["questions"] = [qs[0], second]
    state = interview_state(t, [], [], parse_guide(GUIDE))
    ex = state["cells"][qs[0]["guide_code"]]["exchanges"][0]
    assert ex["dialogue"], "the question's own turn must be shown"
