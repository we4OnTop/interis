"""Word edits, review decisions, extracts, exports and the workflow page (contract section 7)."""

import io
import json
import sys
import time

import pytest
from fastapi.testclient import TestClient

from interis.analysis.analyze import AnalysisOptions, analyze
from interis.analysis.guide import parse_guide
from interis.config import Paths
from interis.web import extracts as export
from interis.web.app import create_app
from interis.web.edits import edits_digest
from interis.web.jobs import JobRunner, job_message
from interis.web.review import _dialogue, passage
from interis.web.store import Store
from interis.web.workflow import STEPS, workflow_state
from tests.test_analysis import FakeEncoder, _interview

GUIDE = ("- F1: Wie sieht Ihr Arbeitsalltag aus?\n- F2: Welche Rolle spielt KI?\n"
         "- F3: Wie gehen Sie mit vertraulichen Daten um?\n- F4: Was wünschen Sie sich?\n")
BASE = "http://127.0.0.1:8765"
TOKEN = "test-token"  # noqa: S105 – fixed token for tests only
H = {"X-Interis": "1"}


@pytest.fixture
def data_dir(tmp_path):
    paths = Paths(tmp_path)
    paths.ensure()
    (tmp_path / "leitfaden.md").write_text(GUIDE, encoding="utf-8")
    t = _interview()
    t.meta.update({"interview_id": "T1", "audio": {"duration_s": 30.0, "sha256": "x"}})
    t.analysis = analyze(t, AnalysisOptions(guide=parse_guide(GUIDE), match_threshold=0.9,
                                            answer_z=0.5), None, FakeEncoder)
    (tmp_path / "exports" / "T1").mkdir(parents=True)
    (tmp_path / "exports" / "T1" / "T1.json").write_text(json.dumps(t.to_dict()),
                                                          encoding="utf-8")
    return paths


@pytest.fixture
def client(data_dir):
    c = TestClient(create_app(data_dir, TOKEN, 8765), base_url=BASE)
    assert c.post("/api/login", json={"token": TOKEN}, headers=H).status_code == 200
    return c


def _store(data_dir):
    return Store(data_dir.root / "interis.db")


def _interview_json(client):
    return client.get("/api/interviews/T1").json()


def _cell(client, code):
    return client.get("/api/projects/1/compare").json()["cells"]["T1"][code]


def _edit(client, **body):
    return client.post("/api/interviews/T1/edits", headers=H, json=body)


def _span(**over):
    return {"turn": 1, "first": 1, "last": 1, "action": "replace", "kind": "correction",
            "text": "war", "tag": "", **over}


# ------------------------------------------------------------------ edits

def test_edit_roundtrip_keeps_the_raw_transcript(client, data_dir):
    assert _edit(client, **_span()).json() == {"ok": True, "edited": 1}
    word = _interview_json(client)["turns"][1]["words"][1]
    assert word["t"].strip() == "war" and word["k"] == "correction" and "g" not in word
    assert word["o"].strip() == "bin"

    raw = json.loads((data_dir.exports / "T1" / "T1.json").read_text(encoding="utf-8"))
    assert raw["turns"][1]["words"][1]["text"].strip() == "bin"

    assert client.post("/api/interviews/T1/edits/revert", headers=H,
                       json={"turn": 1, "first": 0, "last": 3}).status_code == 200
    assert "o" not in _interview_json(client)["turns"][1]["words"][1]


def test_compare_uses_the_effective_text(client):
    _edit(client, turn=3, first=1, last=1, action="replace", kind="correction",
          text="verwenden")
    assert "verwenden" in json.dumps(_cell(client, "F2")["exchanges"][0]["dialogue"])


def test_smoothing_delete_and_replace_with_tags(client):
    assert _edit(client, turn=1, first=2, last=2, action="delete", kind="smoothing",
                 text="", tag="Füllwort").status_code == 200
    assert _edit(client, turn=1, first=3, last=4, action="replace", kind="smoothing",
                 text="Projektleiterin", tag="Grammatik").status_code == 200
    words = _interview_json(client)["turns"][1]["words"]
    assert words[2]["t"] == "" and words[2]["k"] == "smoothing" and words[2]["g"] == "Füllwort"
    assert words[3]["t"].strip() == "Projektleiterin" and words[3]["g"] == "Grammatik"
    assert words[4]["t"] == "" and words[4]["g"] == "Grammatik"
    assert len(_interview_json(client)["edits"]) == 3


def test_invalid_span_is_422(client):
    assert _edit(client, **_span(turn=99)).status_code == 422
    assert _edit(client, **_span(first=999, last=999)).status_code == 422
    assert _edit(client, **_span(first=2, last=1)).status_code == 422
    assert client.post("/api/interviews/T1/edits/revert", headers=H,
                       json={"turn": 1, "first": 5, "last": 1}).status_code == 422


def test_smoothing_needs_a_known_tag(client):
    assert _edit(client, **_span(kind="smoothing", tag="")).status_code == 422
    assert _edit(client, **_span(kind="smoothing", tag="Quatsch")).status_code == 422


def test_correction_rules(client):
    assert _edit(client, **_span(tag="Füllwort")).status_code == 422
    assert _edit(client, **_span(action="delete", text="")).status_code == 422
    assert _edit(client, **_span(text="   ")).status_code == 422


def test_text_rules(client):
    assert _edit(client, **_span(text="a\tb")).status_code == 422
    assert _edit(client, **_span(text="x" * 201)).status_code == 422
    assert _edit(client, **_span(text="ok", tag="Füllwort\x07")).status_code == 422


def test_edit_without_header_is_forbidden(client):
    r = client.post("/api/interviews/T1/edits", json=_span())
    assert r.status_code == 403


def test_edits_do_not_queue_analysis(client, data_dir):
    _edit(client, **_span())
    assert _store(data_dir).jobs(["T1"]) == []


def test_edits_stale_flips_after_analysis_with_digest(client, data_dir):
    assert _interview_json(client)["edits_stale"] is False
    _edit(client, **_span())
    assert _interview_json(client)["edits_stale"] is True

    # what the analysis job does: store the digest of the edits it ran with
    path = data_dir.exports / "T1" / "T1.json"
    raw = json.loads(path.read_text(encoding="utf-8"))
    raw["analysis"]["edits_digest"] = edits_digest(_store(data_dir).word_edits("T1"))
    path.write_text(json.dumps(raw), encoding="utf-8")
    assert _interview_json(client)["edits_stale"] is False

    _edit(client, **_span(text="waren"))
    assert _interview_json(client)["edits_stale"] is True


def test_analyze_endpoint_queues_one_job(client, data_dir):
    assert client.post("/api/interviews/T1/analyze", headers=H).json() == {"queued": 1}
    assert client.post("/api/interviews/T1/analyze", headers=H).json() == {"queued": 0}
    assert [j["kind"] for j in _store(data_dir).jobs(["T1"])] == ["analyze"]


def test_reviewed_flag(client):
    assert _interview_json(client)["reviewed"] is False
    assert client.put("/api/interviews/T1/reviewed", headers=H,
                      json={"reviewed": True}).status_code == 200
    assert _interview_json(client)["reviewed"] is True


def test_retranscription_confirmation_counts_new_markings(client, data_dir):
    audio = data_dir.root / "orig.wav"
    audio.write_bytes(b"RIFF")
    _store(data_dir).register_interview("T1", audio)
    _edit(client, **_span())
    assert client.post("/api/interviews/T1/transcribe", headers=H,
                       json={}).status_code == 409
    assert client.post("/api/interviews/T1/transcribe", headers=H,
                       json={"discard_markings": True}).status_code == 200
    # kept until the new transcript exists: the finished job removes them
    assert len(_store(data_dir).word_edits("T1")) == 1
    assert _edit(client, **_span(text="waren")).status_code == 409  # job queued: no marking


# ------------------------------------------------------------------ decisions

def test_decision_set_clear_and_invalid_code(client):
    url = "/api/interviews/T1/questions/F4/decision"
    assert _cell(client, "F4")["status"] == "missing"
    r = client.put(url, headers=H, json={"reason": "not_relevant", "note": "passt nicht"})
    assert r.status_code == 200
    cell = _cell(client, "F4")
    assert cell["status"] == "explained"
    assert cell["decision"] == {"reason": "not_relevant", "note": "passt nicht"}

    assert client.put("/api/interviews/T1/questions/F9/decision", headers=H,
                      json={"reason": "other"}).status_code == 422
    assert client.put(url, headers=H, json={"reason": "bogus"}).status_code == 422
    assert client.put(url, headers=H, json={"reason": "other", "note": "a\tb"}
                      ).status_code == 422

    assert client.put(url, headers=H, json={"reason": None}).status_code == 200
    assert _cell(client, "F4")["status"] == "missing"
    assert _cell(client, "F4")["decision"] is None


def test_interview_lists_decisions(client):
    client.put("/api/interviews/T1/questions/F3/decision", headers=H,
               json={"reason": "not_asked", "note": ""})
    assert _interview_json(client)["decisions"] == [
        {"guide_code": "F3", "reason": "not_asked", "note": "", "in_guide": True}]


# ------------------------------------------------------------------ extracts

def _extract(client, **over):
    body = {"interview": "T1", "turn": 1, "first": 0, "last": 2, "guide_code": "F1",
            "paraphrase": "Arbeitet als Projektleiterin", **over}
    return client.post("/api/extracts", headers=H, json=body)


def test_extracts_crud(client):
    eid = _extract(client).json()["id"]
    (row,) = client.get("/api/projects/1/extracts").json()["extracts"]
    assert row["id"] == eid and row["interview"] == "T1" and row["guide_code"] == "F1"
    assert row["text"].startswith("Ich bin") and row["start"] < row["end"]
    assert row["paraphrase"] == "Arbeitet als Projektleiterin"

    assert client.patch(f"/api/extracts/{eid}", headers=H,
                        json={"paraphrase": "Projektleiterin"}).status_code == 200
    assert client.get("/api/projects/1/extracts").json()["extracts"][0]["paraphrase"] == \
        "Projektleiterin"

    assert client.delete(f"/api/extracts/{eid}", headers=H).status_code == 200
    assert client.get("/api/projects/1/extracts").json()["extracts"] == []
    assert client.delete(f"/api/extracts/{eid}", headers=H).status_code == 404
    assert client.patch(f"/api/extracts/{eid}", headers=H,
                        json={"paraphrase": "x"}).status_code == 404


def test_extract_validation(client):
    assert _extract(client, turn=99).status_code == 422
    assert _extract(client, first=2, last=1).status_code == 422
    assert _extract(client, guide_code="F9").status_code == 422
    assert _extract(client, paraphrase="   ").status_code == 422
    assert _extract(client, paraphrase="x" * 2001).status_code == 422


def test_export_docx_starts_with_zip_header_and_has_bold_header_row(client):
    from docx import Document

    _extract(client)
    r = client.get("/api/projects/1/extracts/export", params={"format": "docx"})
    assert r.status_code == 200 and r.content.startswith(b"PK")
    assert r.headers["content-disposition"] == 'attachment; filename="extraktion-p1.docx"'
    table = Document(io.BytesIO(r.content)).tables[0]
    assert [c.text for c in table.rows[0].cells] == export.HEADER
    assert all(run.bold for run in table.rows[0].cells[0].paragraphs[0].runs)
    assert len(table.rows) == 2


def test_export_csv_bom_semicolon_and_formula_neutralised(client):
    _extract(client, paraphrase="=cmd Kernaussage")
    r = client.get("/api/projects/1/extracts/export", params={"format": "csv"})
    assert r.content.startswith(b"\xef\xbb\xbf")
    text = r.content.decode("utf-8-sig")
    lines = text.split("\r\n")
    assert lines[0] == ";".join(export.HEADER)
    assert "'=cmd Kernaussage" in lines[1] and lines[1].startswith("T1;F1;")
    assert r.headers["content-disposition"] == 'attachment; filename="extraktion-p1.csv"'


def test_export_rejects_unknown_format(client):
    assert client.get("/api/projects/1/extracts/export",
                      params={"format": "xlsx"}).status_code == 422


def test_csv_neutralises_every_formula_start():
    rows = [export.HEADER, ["=1", "+2", "-3", "@4", "\t5", "ok", "a=b", "", "7"]]
    cells = export.csv_bytes(rows).decode("utf-8-sig").split("\r\n")[1]
    assert cells == "'=1;'+2;'-3;'@4;'\t5;ok;a=b;;7"
    assert export.safe_cell("\r6") == "'\r6" and export.safe_cell("6") == "6"


# ------------------------------------------------------------------ workflow

def test_workflow_endpoint_counts_per_interview(client):
    wf = client.get("/api/projects/1/workflow").json()
    assert [s["id"] for s in wf["steps"]] == [
        "guide", "transcribe", "correct", "smooth", "assign", "explain", "extract", "export"]
    assert wf["steps"] == STEPS and wf["guide_questions"] == 4
    row = wf["interviews"][0]
    assert row["id"] == "T1" and row["transcribed"] is True and row["reviewed"] is False
    assert row["asked"] + row["answered_elsewhere"] + row["omitted"] + row["explained"] \
        + len(row["missing"]) == 4
    assert row["done"]["transcribe"] is True and row["done"]["extract"] is False

    _edit(client, **_span())
    _edit(client, turn=1, first=2, last=2, action="delete", kind="smoothing", text="",
          tag="Füllwort")
    client.put("/api/interviews/T1/reviewed", headers=H, json={"reviewed": True})
    _extract(client)
    row = client.get("/api/projects/1/workflow").json()["interviews"][0]
    assert (row["corrections"], row["smoothing"], row["extracts"]) == (1, 1, 1)
    assert row["edits_stale"] is True and row["reviewed"] is True
    assert row["done"]["correct"] and row["done"]["smooth"] and row["done"]["extract"]


def test_explain_is_done_when_every_missing_question_has_a_reason(client):
    missing = client.get("/api/projects/1/workflow").json()["interviews"][0]["missing"]
    assert missing
    for code in missing:
        client.put(f"/api/interviews/T1/questions/{code}/decision", headers=H,
                   json={"reason": "other", "note": ""})
    row = client.get("/api/projects/1/workflow").json()["interviews"][0]
    assert row["missing"] == [] and row["explained"] == len(missing)
    assert row["done"]["explain"] is True


def test_workflow_state_rules_without_io():
    cells = {"F1": "asked", "F2": "missing"}
    row = {"id": "I01", "transcribed": True, "reviewed": False,
           "edits": [{"kind": "smoothing"}], "edits_stale": False, "cells": cells,
           "unassigned": 2, "extracts": 0}
    out = workflow_state(["F1", "F2"], [row])
    assert out["interviews"][0]["missing"] == ["F2"]
    assert out["interviews"][0]["done"] == {
        "transcribe": True, "correct": False, "smooth": True, "assign": False,
        "explain": False, "extract": False}
    untranscribed = workflow_state(["F1"], [{**row, "transcribed": False, "edits": []}])
    assert untranscribed["interviews"][0]["done"]["assign"] is False


# ------------------------------------------------------------------ projects

def test_project_smoothing_tags(client):
    assert client.get("/api/projects/1").json()["tags"][0] == "Füllwort"
    assert client.patch("/api/projects/1", headers=H,
                        json={"smoothing_tags": "Füllwort\nDialekt"}).status_code == 200
    assert client.get("/api/projects/1").json()["tags"] == ["Füllwort", "Dialekt"]
    too_many = "\n".join(f"Tag {i}" for i in range(31))
    assert client.patch("/api/projects/1", headers=H,
                        json={"smoothing_tags": too_many}).status_code == 422
    assert client.patch("/api/projects/1", headers=H,
                        json={"smoothing_tags": "x" * 41}).status_code == 422
    assert client.patch("/api/projects/1", headers=H,
                        json={"smoothing_tags": "a\tb"}).status_code == 422
    assert client.patch("/api/projects/1", headers=H,
                        json={"smoothing_tags": ""}).status_code == 200
    assert len(client.get("/api/projects/1").json()["tags"]) == 7


# ------------------------------------------------------------------ analysis with edits

def test_cli_analyze_with_edits_keeps_raw_words(data_dir, monkeypatch, tmp_path):
    from interis.cli import build_parser, cmd_analyze
    from interis.pipeline import run
    from interis.web.edits import edits_digest as digest

    def fake_run_analysis(transcript, paths, aopts, speaker_embeddings=None, threads=None):
        transcript.analysis = analyze(transcript, AnalysisOptions(
            guide=parse_guide(GUIDE), match_threshold=0.9, answer_z=0.5), None, FakeEncoder)

    monkeypatch.setattr(run, "run_analysis", fake_run_analysis)
    src = data_dir.exports / "T1" / "T1.json"
    before = json.loads(src.read_text(encoding="utf-8"))
    edits = [{"turn": 0, "word": 2, "action": "delete", "kind": "smoothing", "text": "",
              "tag": "Füllwort"}]
    edits_file = tmp_path / "edits.json"
    edits_file.write_text(json.dumps(edits), encoding="utf-8")

    args = build_parser().parse_args(["--data-dir", str(data_dir.root), "analyze", str(src),
                                      "--edits", str(edits_file), "--formats", "json"])
    assert cmd_analyze(args, data_dir) == 0

    after = json.loads(src.read_text(encoding="utf-8"))
    assert after["turns"] == before["turns"], "raw words must never change"
    assert after["analysis"]["edits_digest"] == digest(edits)
    question = next(q for q in after["analysis"]["questions"] if q["turn"] == 0)
    assert "Ihr" not in question["text"], "the analysis runs on the edited text"


def test_cli_rejects_malformed_edits(data_dir, tmp_path):
    from interis.cli import build_parser, cmd_analyze

    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps([{"turn": "x"}]), encoding="utf-8")
    src = data_dir.exports / "T1" / "T1.json"
    args = build_parser().parse_args(["--data-dir", str(data_dir.root), "analyze", str(src),
                                      "--edits", str(bad)])
    assert cmd_analyze(args, data_dir) == 2


def test_analysis_job_passes_edits_file_and_removes_it(data_dir):
    import sys
    import time

    from interis.web.jobs import JobRunner

    store = _store(data_dir)
    store.set_word_edits("T1", [{"turn": 1, "word": 1, "action": "replace",
                                 "kind": "correction", "text": "war", "tag": ""}])
    script = ("import json,sys; p=sys.argv[sys.argv.index('--edits')+1];"
              "print(len(json.load(open(p, encoding='utf-8'))), flush=True)")

    class Fake(JobRunner):
        def command(self, job):
            return [sys.executable, "-c", script]

    runner = Fake(data_dir, store, lambda _i: None, lambda _i: "")
    job = store.add_job("analyze", "T1")
    runner.start()
    try:
        deadline = time.time() + 30
        while store.job(job)["status"] not in ("done", "failed") and time.time() < deadline:
            time.sleep(0.1)
    finally:
        runner.stop()
    assert store.job(job)["status"] == "done" and store.job(job)["message"] == "1"
    assert list(data_dir.tmp.glob("edits-*.json")) == []


# ------------------------------------------------------------------ one kind per word

def test_one_kind_of_edit_per_word(client):
    assert _edit(client, turn=1, first=2, last=2, action="delete", kind="smoothing",
                 text="", tag="Füllwort").status_code == 200
    # a correction over a smoothed word is refused; the reason is not silently replaced
    r = _edit(client, **_span(first=1, last=2))
    assert r.status_code == 409
    assert r.json()["detail"] == "Diese Stelle hat schon eine andere Änderung: erst zurücknehmen"
    assert _interview_json(client)["turns"][1]["words"][2]["g"] == "Füllwort"
    # the same kind overwrites
    assert _edit(client, turn=1, first=2, last=2, action="delete", kind="smoothing",
                 text="", tag="Dialekt").status_code == 200
    assert _interview_json(client)["turns"][1]["words"][2]["g"] == "Dialekt"


# ------------------------------------------------------------------ position-keyed writes

LINK = {"interview": "T1", "turn": 1, "first": 0, "last": 2, "guide_code": "F4"}


def test_position_writes_wait_for_a_queued_transcription(client, data_dir):
    audio = data_dir.root / "orig.wav"
    audio.write_bytes(b"RIFF")
    store = _store(data_dir)
    store.register_interview("T1", audio)
    link_id = client.post("/api/links", headers=H, json=LINK).json()["id"]
    eid = _extract(client).json()["id"]
    url = "/api/interviews/T1/questions/F3/decision"
    assert client.post("/api/interviews/T1/transcribe", headers=H,
                       json={"discard_markings": True}).status_code == 200

    assert client.post("/api/questions", headers=H, json={
        "interview": "T1", "turn": 2, "first": 0, "last": 1}).status_code == 409
    assert client.post("/api/questions/reset", headers=H, json={
        "interview": "T1", "turn": 1, "first": 0}).status_code == 409
    assert client.post("/api/links", headers=H, json=LINK).status_code == 409
    assert client.patch(f"/api/links/{link_id}", headers=H,
                        json={"omitted": True}).status_code == 409
    assert client.delete(f"/api/links/{link_id}", headers=H).status_code == 409
    assert _extract(client).status_code == 409
    assert client.patch(f"/api/extracts/{eid}", headers=H,
                        json={"paraphrase": "neu"}).status_code == 409
    assert client.delete(f"/api/extracts/{eid}", headers=H).status_code == 409
    assert client.put(url, headers=H, json={"reason": "other"}).status_code == 409
    assert client.put("/api/interviews/T1/reviewed", headers=H,
                      json={"reviewed": True}).status_code == 409
    assert _edit(client, **_span()).status_code == 409
    # nothing was written meanwhile
    assert store.links("T1")[0]["omitted"] is False and store.decisions("T1") == []


def test_relinking_keeps_the_stored_omitted_flag_and_note(client, data_dir):
    client.post("/api/links", headers=H, json={**LINK, "omitted": True,
                                                "note": "Schon in Frage 2 erzählt"})
    client.post("/api/links", headers=H, json=LINK)  # no flag, no note given
    (row,) = _store(data_dir).links("T1")
    assert row["omitted"] is True and row["note"] == "Schon in Frage 2 erzählt"
    client.post("/api/links", headers=H, json={**LINK, "omitted": False, "note": ""})
    (row,) = _store(data_dir).links("T1")
    assert row["omitted"] is False and row["note"] == ""


def test_links_carry_the_effective_passage(client):
    client.post("/api/links", headers=H, json=LINK)
    (lk,) = _interview_json(client)["links"]
    assert lk["text"].startswith("Ich bin") and lk["turn"] == 1
    assert 0 <= lk["start"] < lk["end"]


def test_compare_reports_stale_analysis_per_interview(client):
    assert client.get("/api/projects/1/compare").json()["stale"] == {"T1": False}
    _edit(client, **_span())
    assert client.get("/api/projects/1/compare").json()["stale"] == {"T1": True}


def test_running_analysis_does_not_block_a_new_one(client, data_dir):
    store = _store(data_dir)
    running = store.add_job("analyze", "T1")
    store.update_job(running, status="running")
    assert client.post("/api/interviews/T1/analyze", headers=H).json() == {"queued": 1}
    assert client.post("/api/interviews/T1/analyze", headers=H).json() == {"queued": 0}


# ------------------------------------------------------------------ input rules

@pytest.mark.parametrize("text", ["\u202eabc", "a\u200bb", "\ud800", "a\u2028b", "a\u2029b"])
def test_format_surrogate_and_separator_characters_are_refused(client, text):
    # json.dumps escapes a lone surrogate, so the body stays valid UTF-8
    def raw_post(url, body):
        return client.post(url, headers={**H, "Content-Type": "application/json"},
                           content=json.dumps(body))

    assert raw_post("/api/interviews/T1/edits", _span(text=text)).status_code == 422
    assert raw_post("/api/extracts", {"interview": "T1", "turn": 1, "first": 0, "last": 2,
                                      "guide_code": "F1", "paraphrase": text}
                    ).status_code == 422


def test_multiline_text_keeps_line_breaks(client):
    assert _extract(client, paraphrase="Erste Zeile\nZweite Zeile").status_code == 200


# ------------------------------------------------------------------ workflow and extracts

def test_extracts_of_removed_guide_questions_are_kept_but_not_counted(client):
    _extract(client, guide_code="F1")
    _extract(client, guide_code="F4")
    kept = "".join(line for line in GUIDE.splitlines(True) if "F4" not in line)
    assert client.put("/api/projects/1/guide", headers=H, json={"text": kept}).status_code == 200
    rows = client.get("/api/projects/1/extracts").json()["extracts"]
    assert sorted((r["guide_code"], r["in_guide"]) for r in rows) == [("F1", True), ("F4", False)]
    row = client.get("/api/projects/1/workflow").json()["interviews"][0]
    assert row["extracts"] == 1 and row["done"]["extract"] is True
    csv_text = client.get("/api/projects/1/extracts/export", params={"format": "csv"}
                          ).content.decode("utf-8-sig")
    assert "T1;F4;" in csv_text  # the export keeps every row


def test_extract_outside_the_transcript_is_neither_listed_nor_counted(client, data_dir):
    _store(data_dir).add_extract("T1", "F1", 99, 0, 0, "verwaist")
    assert client.get("/api/projects/1/extracts").json()["extracts"] == []
    row = client.get("/api/projects/1/workflow").json()["interviews"][0]
    assert row["extracts"] == 0 and row["done"]["extract"] is False


def test_exports_survive_control_characters_in_the_guide(client):
    from docx import Document

    text = GUIDE.replace("Wie sieht", "Wie\x1b sieht")
    assert client.put("/api/projects/1/guide", headers=H, json={"text": text}).status_code == 200
    _extract(client)
    r = client.get("/api/projects/1/extracts/export", params={"format": "docx"})
    assert r.status_code == 200
    assert Document(io.BytesIO(r.content)).tables[0].rows[1].cells[2].text == \
        "Wie sieht Ihr Arbeitsalltag aus?"
    assert client.get("/api/projects/1/extracts/export",
                      params={"format": "csv"}).status_code == 200


def test_saving_the_guide_writes_the_codes_of_unlabelled_questions(client):
    text = GUIDE + "- Neue Frage ohne Code?\n"
    assert client.put("/api/projects/1/guide", headers=H, json={"text": text}).status_code == 200
    saved = client.get("/api/projects/1").json()["guide_text"]
    assert saved == GUIDE + "- F5: Neue Frage ohne Code?\n"
    assert [q["code"] for q in client.get("/api/projects/1").json()["guide"]["questions"]] == \
        ["F1", "F2", "F3", "F4", "F5"]


def test_cli_transcription_replaces_the_markings_with_the_transcript(data_dir, monkeypatch,
                                                                     tmp_path):
    from interis.cli import build_parser, cmd_transcribe
    from interis.pipeline import run

    store = _store(data_dir)
    store.set_word_edits("T1", [{"turn": 1, "word": 1, "action": "replace",
                                 "kind": "correction", "text": "war", "tag": ""}])
    store.set_reviewed("T1", True)
    audio = tmp_path / "rec.wav"
    audio.write_bytes(b"RIFF")
    new = _interview()
    new.meta.update({"interview_id": "T1", "audio": {"duration_s": 30.0, "sha256": "y"}})
    monkeypatch.setattr(run, "run_pipeline", lambda *_a, **_k: new)

    args = build_parser().parse_args(["--data-dir", str(data_dir.root), "transcribe",
                                      str(audio), "--id", "T1", "--formats", "json"])
    assert cmd_transcribe(args, data_dir) == 0
    assert store.word_edits("T1") == [] and store.reviewed("T1") is False


def test_workflow_explain_needs_a_guide():
    row = {"id": "I01", "transcribed": True, "reviewed": False, "edits": [], "edits_stale": False,
           "cells": {}, "unassigned": 0, "extracts": 0}
    out = workflow_state([], [row])
    assert out["interviews"][0]["done"]["explain"] is False


def test_overlapping_question_marks_are_shown_once():
    from interis.analysis.analyze import AskedQuestion
    from tests.test_analysis import _turn as make_turn

    t = _interview()
    t.turns = [make_turn("S0", 0, "Wie war Ihre Kindheit?")]

    def question(first, last):
        p = passage(t, 0, first, last)
        return AskedQuestion(id="", turn=0, first=first, last=last, start=p["start"],
                             end=p["end"], text=p["text"], speaker="S0", kind="question",
                             reasons=["manual"], confidence=1.0, match="main", guide_code="F1")

    pieces = _dialogue(t, 0, 1, [question(0, 3), question(2, 3)])[0]["pieces"]
    assert [p["text"] for p in pieces] == ["Wie war Ihre Kindheit?"]


# ------------------------------------------------------------------ transcription jobs

def _run_queue(data_dir, script, job_ids):
    """Run the job queue with ``script`` as the child process until the jobs have ended."""
    store = _store(data_dir)

    class Fake(JobRunner):
        def command(self, job):
            return [sys.executable, "-c", script]

    runner = Fake(data_dir, store, lambda _i: None, lambda _i: "")
    runner.start()
    try:
        deadline = time.time() + 30
        while time.time() < deadline and any(
                store.job(j)["status"] not in ("done", "failed", "cancelled") for j in job_ids):
            time.sleep(0.05)
    finally:
        runner.stop()
    return [store.job(j) for j in job_ids]


def _retranscribe(client, data_dir):
    audio = data_dir.root / "orig.wav"
    audio.write_bytes(b"RIFF")
    _store(data_dir).register_interview("T1", audio)
    _edit(client, **_span())
    client.put("/api/interviews/T1/reviewed", headers=H, json={"reviewed": True})
    r = client.post("/api/interviews/T1/transcribe", headers=H, json={"discard_markings": True})
    assert r.status_code == 200
    return r.json()["job"]


def test_finished_transcription_removes_the_markings(client, data_dir):
    job = _retranscribe(client, data_dir)
    (done,) = _run_queue(data_dir, "print('ok')", [job])
    assert done["status"] == "done"
    assert _store(data_dir).word_edits("T1") == [] and not _store(data_dir).reviewed("T1")


def test_failed_transcription_keeps_the_markings(client, data_dir):
    job = _retranscribe(client, data_dir)
    script = "import sys; print('ERROR: boom', file=sys.stderr); sys.exit(1)"
    (failed,) = _run_queue(data_dir, script, [job])
    assert failed["status"] == "failed" and failed["message"] == "ERROR: boom"
    assert len(_store(data_dir).word_edits("T1")) == 1


def test_job_cancelled_before_it_starts_stays_cancelled(data_dir):
    store = _store(data_dir)
    job_id = store.add_job("analyze", "T1")
    runner = JobRunner(data_dir, store, lambda _i: None, lambda _i: "")
    job = store.job(job_id)  # the worker read the job as queued ...
    assert runner.cancel(job_id)  # ... then it was cancelled
    runner._execute(job, [sys.executable, "-c", "print('ran')"])
    assert store.job(job_id)["status"] == "cancelled"


def test_job_messages_show_no_folders():
    text = "[Errno 13] Permission denied: '/data/tmp/edits-7.json'\nmore"
    assert job_message(text) == "[Errno 13] Permission denied: 'edits-7.json'"
    assert job_message("ERROR: file not found: C:\\Daten\\x.wav") == \
        "ERROR: file not found: x.wav"
