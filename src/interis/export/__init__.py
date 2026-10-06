"""Exporters. JSON is lossless (words, timings, probabilities, analysis, metadata); TXT and
DOCX are reading/review formats. Transcription rules (e.g. Dresing & Pehl) will be added
later as separate transformations applied at export time – the stored data stays raw."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from interis.pipeline.types import Transcript

LOW_CONFIDENCE = 0.5  # words below this Whisper probability are highlighted for review
_SUGGESTION_LABEL = {"anticipated": "vorweg", "later": "später", "unasked": "ohne Frage"}


def fmt_time(seconds: float) -> str:
    s = int(seconds)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def speaker_name(transcript: Transcript, label: str | None) -> str:
    if label is None:
        return "?"
    for sp in transcript.speakers:
        if sp["label"] == label:
            return sp.get("display_name") or label
    return label


def question_tag(q: dict[str, Any]) -> str:
    if q.get("match") == "main":
        return f"[{q['guide_code']}]"
    if q.get("match") == "probe":
        return f"[{q['guide_code']} Nachfrage]"
    if q.get("match") == "followup":
        return "[Nachfrage]"
    return "[Impuls]" if q.get("kind") == "prompt" else "[Frage]"


def _questions_by_turn(transcript: Transcript) -> dict[int, list[dict[str, Any]]]:
    out: dict[int, list[dict[str, Any]]] = {}
    for q in transcript.analysis.get("questions", []):
        out.setdefault(q["turn"], []).append(q)
    return out


def write_json(transcript: Transcript, path: Path) -> None:
    path.write_text(json.dumps(transcript.to_dict(), ensure_ascii=False, indent=2),
                    encoding="utf-8")


def write_txt(transcript: Transcript, path: Path) -> None:
    lines = [f"# {transcript.meta['interview_id']} – {transcript.meta['note']}", ""]
    by_turn = _questions_by_turn(transcript)
    for ti, turn in enumerate(transcript.turns):
        starts = {q["first"]: question_tag(q) for q in by_turn.get(ti, [])}
        parts = []
        for wi, w in enumerate(turn.words):
            if wi in starts:
                parts.append(f" {starts[wi]}")
            parts.append(w.text)
        text = "".join(parts).strip()
        lines += [f"[{fmt_time(turn.start)}] {speaker_name(transcript, turn.speaker)}: {text}",
                  ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def coverage(transcript: Transcript) -> list[dict[str, Any]]:
    """Per guide question: when it was asked, direct answer, other possible answers."""
    analysis = transcript.analysis
    guide = analysis.get("guide")
    if not guide:
        return []
    questions = analysis.get("questions", [])
    answers = {a["question_id"]: a for a in analysis.get("answers", [])}
    rows = []
    for gq in guide["questions"]:
        asked = [q for q in questions if q.get("guide_code") == gq["code"]
                 and q.get("match") == "main"]
        direct = [p for q in asked for p in answers.get(q["id"], {}).get("passages", [])]
        elsewhere = [s for s in analysis.get("suggestions", []) if s["guide_code"] == gq["code"]]
        rows.append({"code": gq["code"], "text": gq["text"], "asked": asked,
                     "direct": direct, "elsewhere": elsewhere})
    return rows


def write_docx(transcript: Transcript, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor

    blue = RGBColor(0x1F, 0x4E, 0xA8)
    meta = transcript.meta
    doc = Document()
    doc.add_heading(f"Transkript {meta['interview_id']}", level=1)

    info = doc.add_table(rows=0, cols=2)
    roles = transcript.analysis.get("roles", {})
    rows = [
        ("Status", "Rohtranskript (maschinell, ungeprüft)"),
        ("Dauer", fmt_time(meta["audio"]["duration_s"])),
        ("Erstellt", meta["created_at"]),
        ("Modelle", ", ".join(f"{m['repo']}@{m['revision'][:10]}"
                              for m in meta["models"].values())),
        ("Audio SHA-256", meta["audio"]["sha256"]),
    ]
    if roles.get("method") and roles["method"] != "none":
        method = {"voice": "Stimmprofil", "heuristic": "Heuristik (Fragenanteil)"}
        rows.append(("Rollenzuordnung", method.get(roles["method"], roles["method"])))
    for key, value in rows:
        cells = info.add_row().cells
        cells[0].text, cells[1].text = key, value

    legend = doc.add_paragraph()
    legend.add_run("Legende: ").bold = True
    q_run = legend.add_run("blau [F1]")
    q_run.bold, q_run.font.color.rgb = True, blue
    legend.add_run(" = erkannte Frage mit Leitfaden-Code, ")
    marked = legend.add_run("gelb")
    marked.font.highlight_color = WD_COLOR_INDEX.YELLOW
    legend.add_run(f" = unsicher erkannt (< {LOW_CONFIDENCE:.0%}), ")
    over = legend.add_run("kursiv")
    over.italic = True
    legend.add_run(" = gleichzeitiges Sprechen.")

    by_turn = _questions_by_turn(transcript)
    for ti, turn in enumerate(transcript.turns):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        head = p.add_run(f"[{fmt_time(turn.start)}] {speaker_name(transcript, turn.speaker)}: ")
        head.bold = True
        spans = by_turn.get(ti, [])
        for wi, w in enumerate(turn.words):
            q = next((q for q in spans if q["first"] <= wi <= q["last"]), None)
            if q is not None and wi == q["first"]:
                tag = p.add_run(("" if wi == 0 else " ") + question_tag(q))
                tag.bold, tag.font.color.rgb = True, blue
            text = w.text.lstrip() if wi == 0 else w.text
            run = p.add_run(text)
            if q is not None:
                run.bold, run.font.color.rgb = True, blue
            if w.prob < LOW_CONFIDENCE:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if w.overlap:
                run.italic = True

    rows_cov = coverage(transcript)
    if rows_cov:
        doc.add_heading("Leitfaden-Abdeckung (Vorschläge, ungeprüft)", level=2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, title in zip(table.rows[0].cells,
                               ("Code", "Leitfadenfrage", "Gestellt", "Evtl. auch beantwortet"),
                               strict=True):
            cell.text = title
        for row in rows_cov:
            cells = table.add_row().cells
            cells[0].text = row["code"]
            cells[1].text = row["text"]
            cells[2].text = (", ".join(fmt_time(q["start"]) for q in row["asked"])
                             or "nicht gestellt")
            cells[3].text = "; ".join(
                f"{fmt_time(s['passages'][0]['start'])} ({_SUGGESTION_LABEL[s['type']]}, "
                f"{s['score']:.2f})" for s in row["elsewhere"]) or "–"
    doc.save(str(path))
