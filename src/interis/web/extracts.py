"""The extraction table (one row per extract) and its exports: Word and CSV (Excel).

Pure functions; the exports are built in memory and never written to disk.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from docx import Document

HEADER = ["Gespräch", "Frage", "Leitfadenfrage", "Kernaussage", "Zitat", "Zeit"]
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CSV_MIME = "text/csv; charset=utf-8"
# Spreadsheet programs run a cell that starts with one of these as a formula.
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def clock(seconds: float) -> str:
    h, rest = divmod(int(seconds), 3600)
    m, s = divmod(rest, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def table(extracts: list[dict[str, Any]], titles: dict[str, str]) -> list[list[str]]:
    """Header plus one row per extract. ``extracts``: rows of GET /extracts;
    ``titles``: guide code -> question text."""
    rows = [[e["interview"], e["guide_code"], titles.get(e["guide_code"], ""),
             e["paraphrase"], e["text"], f"{clock(e['start'])}–{clock(e['end'])}"]
            for e in extracts]
    return [HEADER, *rows]


def csv_bytes(rows: list[list[str]]) -> bytes:
    """UTF-8 with BOM (so Excel reads umlauts), semicolons, CRLF line ends."""
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerows([[safe_cell(c) for c in row] for row in rows])
    return buf.getvalue().encode("utf-8-sig")


def safe_cell(cell: str) -> str:
    """CSV formula injection: a leading quote makes the cell plain text."""
    return "'" + cell if cell.startswith(FORMULA_START) else cell


def docx_bytes(rows: list[list[str]]) -> bytes:
    doc = Document()
    doc.core_properties.title = "Extraktion"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    grid = doc.add_table(rows=0, cols=len(HEADER))
    grid.style = "Table Grid"
    for i, row in enumerate(rows):
        for cell, value in zip(grid.add_row().cells, row, strict=True):
            cell.paragraphs[0].add_run(value).bold = i == 0
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
